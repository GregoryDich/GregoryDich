#!/usr/bin/env python3
"""Build site/MITODERM-3D-story.xlsx from site/story3d.json (the 3D scroll story of every page, 11/10/2026).

story3d.json holds, per page, the story written by the workflow (writer -> marketing and fact critics -> fixer):
screens with the 3D action and text KEYS, the purchase screen, 3D assets, new strings and questions. Texts are never
typed here: each key is resolved against the reviewed sources (storyboard.he.json, product-pages*.json) or the
reviewed new strings, so the workbook shows exactly the Hebrew and English that passed review. An unknown key is
marked, not guessed.

Usage: python3 site/tools/story3d_xlsx.py
"""
import json, os, re
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

SITE = os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(__file__)), '..'))
load = lambda name: json.load(open(os.path.join(SITE, name), encoding='utf-8'))

data = load('story3d.json')
he_pages = {p['slug']: {i['key']: i for i in p['items']} for p in load('storyboard.he.json')['pages']}
cat_en = {s['slug']: s for s in load('product-pages.json')}
cat_he = {s['slug']: h for s, h in zip(load('product-pages.json'), load('product-pages.he.json'))}
nav = he_pages['home']['nav.blocks']
nav_he, nav_en = nav['he'].split(' · '), nav['en'].split(' · ')
new = {(n['slug'], n['key']): n for n in data.get('new_strings', [])}
for h in data.get('hebrew', []):
    new.setdefault((h['slug'], h['key']), {}).update(he=h['he'])

# a new key whose Hebrew already passed review elsewhere: take it verbatim instead of a fresh translation
REVIEWED = {('catalog', 'catalog.chips'): 'הכל · אקסוזומים · פילינג · מסכות · שיער · ביו-ספיקולות · טיפוח ביתי · מכשור'}  # sheet Catalog, Filters · Chips HE
for k, he in REVIEWED.items():
    new.setdefault(k, {}).update(he=he)
MISSING = '⚠ нет ключа: '
PENDING = '⏳ перевод на иврит и две проверки идут'

# The 'who performs the treatment' block (owner 11/10): screen 8 on every page, right after the price screen.
# Texts come from the owner's Make site (sources/make-dictionary-spec.json) and from mitoderm.com (sources/mitoderm-com).
MAKE = load('sources/make-dictionary-spec.json')
FORM = load('sources/mitoderm-com/form-he.json')
FORM_EN = {  # English site form, from the owner's screenshots of mitoderm.com (11/10)
    'form.titleP1': 'Get Personalized', 'form.titleP2': 'Solutions', 'form.titleP3': 'for Your Practice',
    'form.subtitle': 'Our Specialist Will Reach Out',  # site dictionary; owner 11/10: never promise a callback time
    'form.placeholderInputName': 'Name',
    'form.placeholderEmailName': 'E-mail', 'form.placeholderPhoneName': 'Phone Number', 'form.placeholderProfession': 'Profession',
    'form.checkboxLabel': 'I agree with rules of Mitoderm *', 'form.sharing': 'Your information will never be shared with third parties *',
    'buttons.requestCallback': 'Request A Callback', 'form.sent.mobTitle': 'Congratulations!', 'form.sent.text': 'Your form has been submitted',
    'form.sent.textColored': 'successfully', 'buttons.contactForPrice': 'Contact for price', 'specialists.title': 'Certified cosmetologists and clinics',
    'buttons.learnMore': 'More info', 'form.eventTitle': 'Reserve your spot for the Event', 'form.placeholderInputID': 'ID Card Number',
    'form.numberTickets': 'Tickets', 'form.total': 'TOTAL', 'form.promo': 'Promo Code', 'form.apply': 'Apply', 'buttons.reserveSeat': 'Reserve a Seat'}
SITE_HE = load('sources/mitoderm-com/products-he.json')  # mitoderm.com product texts (owner 11/10: very good, not perfect — best of site + catalogue)
SITE_EN = load('sources/mitoderm-com/products-en.json')
CONTACTS = load('sources/mitoderm-com/contacts.json')
TAGS = {'face': ('פנים', 'Face', 'заказчик'), 'scalp': ('קרקפת', 'Scalp', 'заказчик'),
        'trich': (MAKE['prof.trichologist']['he'], MAKE['prof.trichologist']['en'], 'Make')}
SPEC_LINE = {'v-tech-system': 'tech.face.1.body', 'exo-nad': 'tech.face.1.body', 'exocell-mask': 'tech.face.1.body',
             'exosignal-hair': 'tech.hair.1.body', 'mitopen': 'tech.devices.3.body', 'mitoscan': 'tech.devices.3.body',
             'micro-boost-10': 'tech.biospicules.1.body'}
SPEC_TAGS = {'exosignal-hair': ['scalp', 'trich'], 'exosignal-spray': ['scalp', 'trich'], 'mitopen': ['face', 'scalp'],
             'home': ['face', 'scalp', 'trich'], 'catalog': ['face', 'scalp', 'trich']}
# owner's photos in the Figma collage (site/photo-catalog.md), by page
PHOTOS = {'exocell-mask': '11–24 (саше, двойная маска, нанесение)', 'v-tech-system': '40, 42, 44–50, 53, 54, 60, 78–80, 82, 83',
          'exo-nad': '26–32, 35, 39, 51, 63, 64, 84', 'exosignal-hair': '36, 47, 55, 58', 'exosignal-spray': '2, 3, 21, 33, 34, 41, 43, 59, 65',
          'exotech-gel': '37, 38, 52', 'mitopen': '56, 57', 'derma-recovery-cream': '1, 5, 73', 'micro-boost-10': '6, 67–69, 72, 74, 77',
          'cell-renew-1-5': '10, 70, 75, 67–69', 'cellular-age-defense-2-5': '7, 9, 67–69', 'mitoscan': '81',
          'home': '66, 71, 76', 'catalog': 'пакшоты всех средств'}


def form_value(path):
    obj = FORM
    for part in path.split('.'):
        obj = obj[part]
    return obj


def walk(obj, path):
    """cat path like 'steps[0][0]' or 'ingredients' -> value"""
    for part in re.findall(r'[A-Za-z_0-9]+|\[\d+\]', path):
        obj = obj[int(part[1:-1])] if part.startswith('[') else obj[part]
    return obj


def flat(v):
    if isinstance(v, list):
        return ' · '.join(flat(x) for x in v)
    return str(v)


UI = load('proto/exocell-he-ui.json')
TOKEN = (r'«[^»]*»|(?:new:)?(?:[a-z0-9-]+:)?(?:cat|sb)\[[a-z0-9-]+\]\.[A-Za-z0-9_.\[\]-]+'
         r'|(?:new:)?(?:[a-z0-9-]+:)?[a-z][a-z0-9_-]*(?:\.[A-Za-z0-9_-]+|\[\d+\])+(?:#\d+)?|\bnav\[\d\]|\bvp\b|\bnone\b')  # page UI strings (Hebrew), reviewed with the MASK page


def resolve(slug, key):
    """-> (he, en) for a text key, or a marked miss. Accepts the forms the story writers used:
    'a + b' or 'a, b' (several keys), 'page:key' (another page), 'key[i]' (i-th item of a ' · ' list), a trailing '(note)'."""
    if not key:
        return '', ''
    key = key.strip()
    key = re.sub(r'(b\d\.label)(\d)[–-](\d)', lambda m: ' · '.join(m.group(1) + str(i) for i in range(int(m.group(2)), int(m.group(3)) + 1)), key)
    if re.search(r'[\s·+,/«]', key):                             # several keys, maybe mixed with notes
        toks = re.findall(TOKEN, re.sub(r'\([^()]*\)', ' ', key))
        if len(toks) > 1 or (toks and toks[0] != key):
            got = [(t[1:-1], t[1:-1]) if t.startswith('«') else resolve(slug, t) for t in toks]
            got = [g for g in got if g[0].strip(' –—-')]
            return ' · '.join(g[0] for g in got), ' · '.join(g[1] for g in got)
    key = re.sub(r'\s*\([^()]*\)', '', key)                  # writer notes in brackets
    key = re.split(r'\s+(?:—|→)\s*', key)[0].strip()           # '... → screen 7', '... — one button'
    key = ' '.join(w for w in key.split() if not re.search('[а-яё]', w, re.I) and not re.fullmatch(r'[A-Z0-9%-]+', w))  # notes in key fields
    key = re.sub(r'(^|(?<=[\s+·/,]))new:', '', key).strip()      # 'new:key' = a reviewed new string
    if key.lower() in ('none', '-', '—', ''):
        return '', ''
    m = re.fullmatch(r'(cat|sb)\[([a-z0-9-]+)\]\.(.+)', key)   # cat[slug].path / sb[slug].key
    if m:
        return resolve(m.group(2), ('cat.' if m.group(1) == 'cat' else '') + m.group(3))
    m = re.fullmatch(r'(.+)#(\d+)', key)                         # key#n = n-th item (1-based)
    if m:
        return resolve(slug, f'{m.group(1)}[{int(m.group(2)) - 1}]')
    parts = [p for p in re.split(r'\s*[+·]\s*|\s+/\s+|,\s*(?=[a-z])|\s+(?=[a-z][a-z0-9-]*[.:\[])', key) if p]
    if len(parts) > 1:
        got = [resolve(slug, p) for p in parts]
        return ' · '.join(g[0] for g in got if g[0]), ' · '.join(g[1] for g in got if g[1])
    m = re.fullmatch(r'([a-z0-9-]+):(.+)', key)
    if m and not key.startswith(('make.', 'site.', 'tag.', 'cat.', 'mdc.', 'ui.')):
        return resolve(m.group(1), m.group(2))
    m = re.fullmatch(r'(b\d\.[a-z]+\d?|vp)\[(\d+)\]', key)
    if m:
        he, en = resolve(slug, m.group(1))
        if not he.startswith(MISSING):
            i = int(m.group(2))
            hs, es = he.split(' · '), en.split(' · ')
            return (hs[i] if i < len(hs) else he), (es[i] if i < len(es) else en)
    try:
        if key.startswith('ui.'):
            v = UI[key[3:]]
            return v, (v if re.fullmatch(r'[ -~]+', v) else '')
        m = re.fullmatch(r'nav\[(\d)\]', key)
        if m:
            return nav_he[int(m.group(1))], nav_en[int(m.group(1))]
        if key.startswith('make.'):
            e = MAKE[key[5:]]
            return e['he'], e['en']
        if key.startswith('site.'):
            path = key[5:]
            return form_value(path), FORM_EN.get(path, '')
        if key.startswith('mdc.'):
            return flat(walk(SITE_HE, key[4:])), flat(walk(SITE_EN, key[4:]))
        if key.startswith('tag.'):
            t = TAGS[key[4:]]
            return t[0], t[1]
        if key.startswith('cat.'):
            path = key[4:]
            return flat(walk(cat_he[slug], path)), flat(walk(cat_en[slug], path))
        page = he_pages['home'] if key.startswith('home.') else he_pages.get(slug, {})
        if key in page:
            return page[key]['he'], page[key]['en']
        m = re.fullmatch(r'(.+?)\[(\d+)\]', key)
        nk = (slug, m.group(1)) if m and (slug, m.group(1)) in new else None
        if m and not nk and m.group(1).startswith('catalog.') and ('catalog', m.group(1)) in new:
            nk = ('catalog', m.group(1))
        if nk:
            n, i = new[nk], int(m.group(2))
            he, en = n.get('he', PENDING), n.get('en', '')
            hs, es = he.split(' · '), en.split(' · ')
            return (hs[i] if i < len(hs) and he != PENDING else he), (es[i] if i < len(es) else en)
        if key.startswith('catalog.') and ('catalog', key) in new and (slug, key) not in new:
            slug = 'catalog'
        if (slug, key) in new:
            n = new[(slug, key)]
            return n.get('he', PENDING), n.get('en', '')
    except (KeyError, IndexError, TypeError):
        pass
    return MISSING + key, MISSING + key


HEAD = PatternFill('solid', fgColor='1A1817')
HFONT = Font(bold=True, color='FFFFFF')
WRAP = Alignment(wrap_text=True, vertical='top')
RTL = Alignment(wrap_text=True, vertical='top', horizontal='right', readingOrder=2)


def sheet(wb, title, cols, rows, he_cols=()):
    ws = wb.create_sheet(title)
    ws.append([c for c, _ in cols])
    for c in ws[1]:
        c.fill, c.font, c.alignment = HEAD, HFONT, WRAP
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = RTL if ws.cell(1, c.column).value in he_cols else WRAP
    for i, (_, w) in enumerate(cols, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    ws.freeze_panes = 'C2'
    return ws


stories = data['stories']
wb = Workbook(); wb.remove(wb.active)
DRAFT = any('_status' in s for s in stories)
ov_rows = [[s['name'], s.get('_status', 'готово'), s['slug'], s['line'], s['use'], s['colour_env'], s['series_tint_screen'], s['hero_object'],
            s['story_one_line'], ' · '.join(TAGS[t][0] for t in SPEC_TAGS.get(s['slug'], ['face'])), PHOTOS.get(s['slug'], ''),
            len(s.get('new_strings', [])), len(s.get('questions_ru', []))] for s in stories]
ov_cols = [('Страница', 22), ('Статус', 18), ('Slug', 20), ('Линейка', 16), ('Назначение', 14), ('Цветовое окружение', 40),
           ('Экран с оттенком серии', 12), ('Герой-объект', 30), ('Сюжет одной строкой', 60), ('Метки специалистов', 18),
           ('Фото заказчика (№ в коллаже)', 24), ('Новых строк', 10), ('Вопросов', 10)]
if not DRAFT:  # the status column only matters while some pages are still in review
    ov_cols.pop(1); ov_rows = [r[:1] + r[2:] for r in ov_rows]
sheet(wb, 'Overview', ov_cols, ov_rows)
rows = []
for s in stories:
    for sc in s['screens']:
        k, t, l, c = (resolve(s['slug'], sc[x]) for x in ('kicker_key', 'title_key', 'line_key', 'cta_key'))
        labs = [resolve(s['slug'], x) for x in sc['label_keys']]
        rows.append([s['name'], sc['n'], sc['block'], sc['screen_name_ru'], sc['action_3d_ru'], sc['transition_ru'],
                     k[0], t[0], l[0], ' · '.join(x[0] for x in labs), c[0], t[1], l[1], ' · '.join(x[1] for x in labs), c[1], sc['mobile_ru'],
                     ' | '.join([sc['kicker_key'], sc['title_key'], sc['line_key'], ','.join(sc['label_keys']), sc['cta_key']])])
for s in stories:
    tags = SPEC_TAGS.get(s['slug'], ['face'])
    sc = {'n': len(s['screens']) + 1, 'block': '7 Specialists', 'screen_name_ru': 'Кто делает процедуру (после экрана с ценой)',
          'action_3d_ru': 'Объекты уходят, дым линейки успокаивается под тёмной вуалью. Три карточки специалистов: фото сверху плавно '
                          'переходит в текст (ничего не накладывается), имя, город, телефон, метка и «מוסמך», ссылка «קביעת ייעוץ». '
                          'Фильтр «מקצוע» открыт на метке этой страницы; ниже кнопка полного списка и строка для косметологов.',
          'transition_ru': 'Смена экрана без движения объектов; карточки появляются по очереди (0,08 с), как в Make.',
          'kicker_key': 'make.spec.kicker', 'title_key': 'make.spec.title', 'line_key': 'make.' + SPEC_LINE.get(s['slug'], 'spec.intro'),
          'label_keys': ['make.spec.filter.all'] + ['tag.' + t for t in tags] + ['make.spec.certified', 'make.spec.book', 'make.spec.join.title', 'make.spec.join.cta'],
          'cta_key': 'site.specialists.title',
          'mobile_ru': 'Карточки столбиком, фото 228 px с переходом в текст, фильтр-чипы переносятся; липкая панель с ценой внизу.'}
    k, t, l, c = (resolve(s['slug'], sc[x]) for x in ('kicker_key', 'title_key', 'line_key', 'cta_key'))
    labs = [resolve(s['slug'], x) for x in sc['label_keys']]
    rows.append([s['name'], sc['n'], sc['block'], sc['screen_name_ru'], sc['action_3d_ru'], sc['transition_ru'],
                 k[0], t[0], l[0], ' · '.join(x[0] for x in labs), c[0], t[1], l[1], ' · '.join(x[1] for x in labs), c[1], sc['mobile_ru'],
                 ' | '.join([sc['kicker_key'], sc['title_key'], sc['line_key'], ','.join(sc['label_keys']), sc['cta_key']])])
order = {s['name']: i for i, s in enumerate(stories)}
rows.sort(key=lambda r: (order[r[0]], r[1]))
sheet(wb, 'Screens', [('Страница', 20), ('#', 4), ('Блок', 12), ('Экран', 24), ('Что делает 3D', 50), ('Переход на скролле', 40),
                      ('Кикер HE', 18), ('Заголовок HE', 34), ('Фраза HE', 50), ('Метки HE', 40), ('Кнопка HE', 30),
                      ('Title EN', 30), ('Line EN', 46), ('Labels EN', 36), ('CTA EN', 28), ('Телефон', 36), ('Ключи', 40)],
      rows, he_cols=('Кикер HE', 'Заголовок HE', 'Фраза HE', 'Метки HE', 'Кнопка HE'))
rows = []
for s in stories:
    ins = s['inside']
    act, btn = resolve(s['slug'], ins['actives_key']), resolve(s['slug'], ins['button_key'])
    for i, c in enumerate(ins['components']):
        rows.append([s['name'], c['what'], c['qty'], c['volume_dose'], c['source'],
                     act[0] if i == 0 else '', btn[0] if i == 0 else '', 'да' if ins['pro_only'] and i == 0 else ''])
sheet(wb, "What's inside", [('Страница', 20), ('Компонент', 34), ('Кол-во', 10), ('Объём / доза', 26), ('Источник', 30),
                            ('Активы HE', 50), ('Кнопка HE', 28), ('Только для специалистов', 12)], rows, he_cols=('Активы HE', 'Кнопка HE'))
sheet(wb, '3D assets', [('Страница', 20), ('Объект', 34), ('Как строим', 30), ('Текстура / источник', 44), ('Статус', 16)],
      [[s['name'], a['object'], a['build'], a['texture_source'], a['status']] for s in stories for a in s['assets']])
sheet(wb, 'New strings', [('Страница', 20), ('Ключ', 20), ('EN', 50), ('HE', 50), ('Где', 40)],
      [[s['name'], n['key'], n['en'], resolve(s['slug'], n['key'])[0], n['where']] for s in stories for n in s.get('new_strings', [])], he_cols=('HE',))
FORM_ROWS = [
    ('Заявка на средство', 'Заголовок, часть 1', 'site.form.titleP1'), ('Заявка на средство', 'Заголовок, часть 2 (золото)', 'site.form.titleP2'),
    ('Заявка на средство', 'Заголовок, часть 3', 'site.form.titleP3'), ('Заявка на средство', 'Подзаголовок', 'site.form.subtitle'),
    ('Заявка на средство', 'Поле', 'site.form.placeholderInputName'), ('Заявка на средство', 'Поле', 'site.form.placeholderProfession'),
    ('Заявка на средство', 'Поле', 'site.form.placeholderEmailName'), ('Заявка на средство', 'Поле', 'site.form.placeholderPhoneName'),
    ('Заявка на средство', 'Согласие', 'site.form.checkboxLabel'), ('Заявка на средство', 'Кнопка', 'site.buttons.requestCallback'),
    ('Заявка на средство', 'Примечание', 'site.form.sharing'), ('Заявка на средство', 'Загрузка', ''),
    ('Заявка на средство', 'Успех (телефон)', 'site.form.sent.mobTitle'), ('Заявка на средство', 'Успех', 'site.form.sent.text'),
    ('Заявка на средство', 'Успех (золото)', 'site.form.sent.textColored'), ('Заявка на средство', 'Кнопка после успеха', 'site.buttons.learnMore'),
    ('Мероприятие (в конце работ)', 'Заголовок', 'site.form.eventTitle'), ('Мероприятие (в конце работ)', 'Поле', 'site.form.placeholderInputName'),
    ('Мероприятие (в конце работ)', 'Поле', 'site.form.placeholderInputID'), ('Мероприятие (в конце работ)', 'Билеты −/+', 'site.form.numberTickets'),
    ('Мероприятие (в конце работ)', 'Поле', 'site.form.placeholderEmailName'), ('Мероприятие (в конце работ)', 'Поле', 'site.form.placeholderPhoneName'),
    ('Мероприятие (в конце работ)', 'Итого ₪', 'site.form.total'), ('Мероприятие (в конце работ)', 'Промокод', 'site.form.promo'),
    ('Мероприятие (в конце работ)', 'Применить', 'site.form.apply'), ('Мероприятие (в конце работ)', 'Промокод принят', 'site.form.success'),
    ('Мероприятие (в конце работ)', 'Промокод неверный', 'site.form.wrongPromo'), ('Мероприятие (в конце работ)', 'Согласие', 'site.form.checkboxLabel'),
    ('Мероприятие (в конце работ)', 'Кнопка', 'site.buttons.reserveSeat'), ('Мероприятие (в конце работ)', 'Примечание', 'site.form.sharing'),
]
frows = []
for kind, what, key in FORM_ROWS:
    he, en = resolve('', key) if key else ('', '')
    note = 'Индикатор загрузки, поля затемнены (как на сайте)' if what == 'Загрузка' else ''
    frows.append([kind, what, he, en, key or '—', note])
sheet(wb, 'Form', [('Форма', 22), ('Элемент', 24), ('HE', 50), ('EN', 40), ('Ключ', 34), ('Примечание', 40)], frows, he_cols=('HE',))
# mitoderm.com: a very good version of product texts (owner 11/10: take the best of it and of the catalogue); contacts as on the site
MDC = [('V-TECH SYSTEM', ['intro.text', 'fastResult.subtitle', 'fastResult.text', 'solution.main.title', 'solution.main.subtitle']
        + ['solution.main.item1title'] + [f'solution.main.item1textP{i}' for i in range(1, 7)]
        + ['solution.main.item2title'] + [f'solution.main.item2textP{i}' for i in range(1, 6)]
        + ['solution.main.bundleTitle', 'solution.main.bundleText1', 'solution.main.bundleText2', 'solution.main.discountTitle',
           'solution.main.discountLabelText1', 'solution.main.discountLabelText2', 'howToUse.title'] + [f'howToUse.main.item{i}' for i in range(1, 7)]
        + ['synergy.title', 'synergy.item1.title', 'synergy.item1.text1', 'synergy.item1.text2', 'synergy.item1.text3', 'synergy.item2.title',
           'synergy.item2.text', 'synergy.item3.title', 'synergy.item3.text']),
       ('EXOTECH GEL', ['intro.gelSubtitle', 'solution.gel.itemTitle'] + [f'solution.gel.item{i}' for i in range(1, 7)]
        + ['solution.gel.discountTitle', 'solution.gel.discountLabelText2', 'solution.gel.bundleText1', 'solution.gel.bundleText3']
        + [f'howToUse.gel.item{i}' for i in range(1, 6)]),
       ('EXOSIGNAL SPRAY', ['intro.hairSubtitle', 'solution.hair.itemTitle'] + [f'solution.hair.item{i}' for i in range(1, 7)]
        + ['solution.hair.discountTitle', 'solution.hair.discountLabelText2', 'solution.hair.bundleText1', 'solution.hair.bundleText2']),
       ('EXOSIGNAL HAIR', ['intro.signalSubtitle', 'solution.signal.title'] + [f'solution.signal.item{i}' for i in range(1, 7)]
        + ['solution.signal.bundleTitle'] + [f'solution.signal.bundleText{i}' for i in range(1, 5)] + [f'howToUse.hair.item{i}' for i in range(1, 7)]),
       ('Общее', ['unique.textP1', 'unique.textP2', 'unique.textP3', 'buttons.contactForPrice', 'buttons.requestCallback', 'buttons.learnMore'])]
mrows = [[prod, k, *resolve('', 'mdc.' + k)] for prod, keys in MDC for k in keys]
C = CONTACTS
mrows += [['Контакты', 'телефон', f"{C['phone_label_he']} {C['phone']}", C['phone_en']], ['Контакты', 'WhatsApp', C['whatsapp'], C['whatsapp']],
          ['Контакты', 'почта (сайт и формы)', C['email'], C['email']], ['Контакты', 'почта Академии (только мероприятия)', C['email_academy'], C['email_academy']],
          ['Контакты', 'адрес', C['address_he'], C['address_en']],
          ['Контакты', 'часы', f"{C['hours_he']['workWeek']} · {C['hours_he']['friday']}: {C['hours_he']['fridayText']} · {C['hours_he']['saturday']}: {C['hours_he']['closed']}",
           f"{C['hours_en']['workWeek']} · {C['hours_en']['friday']}: {C['hours_en']['fridayText']} · {C['hours_en']['saturday']}: {C['hours_en']['closed']}"],
          ['Контакты', 'срок обратного звонка', 'не указываем (заказчик 11/10)', 'not stated (owner 11/10)']]
sheet(wb, 'mitoderm.com', [('Средство', 18), ('Ключ на сайте', 34), ('HE', 60), ('EN', 60)], mrows, he_cols=('HE',))
# answers already given by the owner (11/10) or found in his sources: mitoderm.com, protocols/src/content.*.mjs
CONTACTS_RU = ('Как на mitoderm.com (заказчик 11/10): кнопка WhatsApp +972-54-762-1889, телефон 054-326-2182, '
               'почта info@mitoderm.com (Mitoderm@gmail.com — только для мероприятий). Срок обратного звонка не пишем.')
SPECIALISTS_RU = 'Таблицу специалистов заказчик пришлёт позже (11/10). До неё кнопка ведёт на форму заявки, без JS — в WhatsApp.'
ANSWERS = [
    ('', ('Контакты', 'контакты', 'WhatsApp ведёт на номер', 'Какие номер WhatsApp'), CONTACTS_RU),
    ('', ('איפה לקנות',), SPECIALISTS_RU),
    ('v-tech-system', ('Флакон или ампула',), 'В тексте — «אמפולות», как на mitoderm.com («5 אמפולות V-Tech Serum בנפח 5 מ׳׳ל כל אחת»); в 3D — флакон, как на фото. Заказчик 11/10: «5 масок и 5 ампул».'),
    ('v-tech-system', ('Колпачок флакона',), 'Серебряный, как на фото 79. Сыворотку набирают в шприц 3 мл иглой для набора (протоколы заказчика, V-TECH + MITOPEN, шаг 1).'),
    ('v-tech-system', ('Гиалуроновая кислота и алоэ',), 'В гель-маске: mitoderm.com, состав V-Tech Gel Mask 5ml — «חומצה היאלורונית במשקל מולקולרי נמוך, אלוורה».'),
    ('v-tech-system', ('Сколько флаконов и саше',), 'Протоколы заказчика: на зону (лицо, шея или декольте) до 2 мл сыворотки и до 2 мл гель-маски; раз в 14 дней, курс 4–6 процедур, затем поддержка раз в 2–3 месяца. Сколько процедур на коробку — в протоколах нет.'),
    ('v-tech-system', ('Шаг 1 протокола',), 'Протокол заказчика называется «V-TECH + MITOPEN»: сыворотку вносят MITOPEN, мезопеном или другим подходящим оборудованием. MITOPEN показываем как перекрёстную продажу.'),
    ('v-tech-system', ('Финиш',), 'Протокол заказчика: рекомендована EXOCELL MASK, в конце DERMA RECOVERY CREAM, защита от солнца обязательна. Даём оба чипа.'),
    ('exo-nad', ('Сколько процедур рассчитано',), 'Протоколы заказчика: набора хватает до 15 процедур, в зависимости от расхода.'),
    ('exocell-mask', ('Слои маски',), 'Заказчик 11/10: маска двухкомпонентная, «2 в 1»; тексты s5.double и kit.double не трогаем. Снимает профессионал в клинике, клиент может использовать после процедуры.'),
    ('exocell-mask', ('времени экспозиции',), 'Протоколы заказчика: оставить маску до 30 минут, затем снять и вбить остаток сыворотки в кожу.'),
    ('exosignal-spray', ('состав спрея',), 'Есть на mitoderm.com: синтетические экзосомы, PDRN низкой молекулярной массы, пантенол, биотин, биомиметические пептиды 8 видов; флакон 15 мл (лист mitoderm.com).'),
    ('exosignal-spray', ('Сколько нажатий',), 'mitoderm.com: раз в день на чистую сухую кожу головы, встряхнуть, распылить у корней, слегка помассировать; лучше вечером, курс два месяца. Число нажатий там не указано.'),
    ('mitoscan', ('Кнопка цены',), 'Да: «לקבלת פרטים על מחירים» — как на mitoderm.com.'),
    ('', ('MITOTECH CELL BOOSTER',), 'В протоколах заказчика: два флакона — бустер и концентрат-порошок, смешать до прозрачного раствора, до 2 мл на лицо, шею и декольте. Линейка там не указана.'),
]
def answer(slug, q):
    for sl, keys, a in ANSWERS:
        if (not sl or sl == slug) and any(k in q for k in keys):
            return a
    return ''
qrows = [[s['name'], q, answer(s['slug'], q)] for s in stories for q in s.get('questions_ru', [])]
sheet(wb, 'Questions', [('Страница', 20), ('Вопрос заказчику', 90), ('Ответ / решение (источник)', 70)], qrows)
# facts checked against the owner's protocols, mitoderm.com and the catalogue after the stories were written (11/10)
RECON = os.path.join(SITE, 'story3d-recon.json')
if os.path.exists(RECON):
    KIND = {'wrong': 'ошибка', 'conflict_between_sources': 'источники расходятся', 'missing_important': 'не хватает важного',
            'answer_to_open_question': 'ответ на вопрос'}
    names = {s['slug']: s['name'] for s in stories}
    rec = sorted(json.load(open(RECON, encoding='utf-8'))['confirmed'], key=lambda f: ([s['slug'] for s in stories].index(f['page']) if f['page'] in names else 99, str(f['screen'])))
    sheet(wb, 'Сверка с источниками', [('Страница', 20), ('Экран', 8), ('Что', 24), ('Тип', 16), ('В сюжете', 44), ('В источнике', 50),
                                       ('Источник', 30), ('Предложение', 50), ('Проверка скептика', 40)],
          [[names.get(f['page'], f['page']), f['screen'], f['field'], KIND.get(f['kind'], f['kind']), f['story_says'], f['source_says'],
            f['source_file'].replace('/home/user/GregoryDich/mitoderm-catalog/', ''), f['proposal_ru'], f.get('verify_ru', '')] for f in rec])
out = os.path.join(SITE, 'MITODERM-3D-story.xlsx')
wb.save(out)
misses = sum(str(c.value).startswith(MISSING) for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value)
print(out, {ws.title: ws.max_row - 1 for ws in wb.worksheets}, 'unresolved keys:', misses)
