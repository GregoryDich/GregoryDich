#!/usr/bin/env python3
"""Hebrew texts for the v2 Story Scroll site.

  python3 he_texts.py inputs <story-final.json> <home-dump.json> <workdir>
      -> site/storyboard.json, site/home.en.json, <workdir>/in/<group>.json
  python3 he_texts.py merge <workdir>/out.json
      -> site/storyboard.he.json (+ coverage / numbers / latin checks)
  python3 he_texts.py xlsx
      -> sheet 'Story HE' in site/mitoderm-site-blocks.xlsx

EXOCELL MASK blocks come from the board (site/review-board-exocell.md,
"Блоки после коллегии"), not from the scenario draft.
"""
import json, os, re, sys

HERE = os.path.dirname(os.path.abspath(__file__))
SITE = os.path.normpath(os.path.join(HERE, '..'))

BLOCK_NAMES = ['Meet', 'Why', 'Inside', 'Feel it', 'Protocol', 'Kit & facts', 'Offer']

# Board version of EXOCELL MASK (review-board-exocell.md, 10/10).
EXOCELL_BOARD = {
    'value_proposition': 'EXOCELL MASK is the soothing finishing step after professional treatments: '
                         'a bio-cellulose second skin that cools, hydrates and supports the barrier '
                         'the moment the procedure ends.',
    'blocks': [
        ('EXOCELL MASK — your second skin',
         'Advanced bio-cellulose with synthetic exosomal structures — the soothing finishing step after professional treatments.',
         ['Post-procedure finish', 'Hydration & cooling', 'Barrier support'],
         "Primary: Contact for price · Secondary: See what's inside ↓"),
        ('Calm after the procedure',
         'After microneedling, peels and spicules, skin asks for hydration and barrier support.',
         ['BEST USED AFTER', 'V-TECH SYSTEM', 'EXO-NAD peel', 'MICRO BOOST 10%', 'Microneedling'],
         'none (chips link to V-TECH SYSTEM, EXO-NAD, MICRO BOOST 10% and MITOPEN pages; the first chip is a kicker)'),
        ('What the sheet holds',
         'Synthetic exosomal structures, polynucleotides, hyaluronic acid and Q10 — kept in continuous contact with the skin by thin bio-cellulose.',
         ['Synthetic exosomal structures', 'Polynucleotides', 'Hyaluronic acid', 'Coenzyme Q10'],
         'Full formula (opens the list of all 8 actives in place)'),
        ('Cool, calm, deeply hydrated',
         'A feeling of comfort and freshness — skin left hydrated, smooth, supple and radiant.',
         ['Intensive hydration', 'Cooling & soothing', 'Less feeling of dryness', 'Smoothing & firming'],
         "Text link 'Contact for price' (desktop)"),
        ('Apply. Absorb. Finish.',
         'On cleansed skin after the procedure — moist, continuous contact with the serum; finish with DERMA RECOVERY CREAM and sun protection.',
         ['1 · Apply · minimal air pockets', '2 · Absorb', '3 · Finish', '+ DERMA RECOVERY CREAM · Sun protection'],
         'none'),
        ('Five masks, one box',
         'Suited to post-procedure protocols — especially dry, tired, sensitive or devitalised skin.',
         ['5 masks per box', 'Single-use', 'Active serum', 'Professional use only'],
         'none'),
        ('Bring EXOCELL MASK to your clinic',
         'For professional use only.',
         ['Training & protocols', 'Exclusive importer of VM Corporation, Italy'],
         'Primary: Contact for price (WhatsApp + form) · Secondary: Add to a V-TECH / EXO-NAD enquiry · Tertiary: phone, email · Mobile sticky bar docks here'),
    ],
}

GROUPS = [
    ('G1-pro-face', ['v-tech-system', 'exo-nad', 'exocell-mask']),
    ('G2-exo-hair-gel', ['exosignal-hair', 'exosignal-spray', 'exotech-gel']),
    ('G3-biospicule', ['micro-boost-10', 'cell-renew-1-5', 'cellular-age-defense-2-5']),
    ('G4-cream-devices', ['derma-recovery-cream', 'mitopen', 'mitoscan']),
    ('G5-knowledge', ['about', 'synthetic-exosomes', 'biospicule']),
    ('G6-home', ['home']),
]


def load(p):
    with open(p, encoding='utf-8') as f:
        return json.load(f)


def dump(obj, p):
    with open(p, 'w', encoding='utf-8') as f:
        json.dump(obj, f, ensure_ascii=False, indent=1)
        f.write('\n')


def page_items(page):
    """Flat list of translatable strings of one story page."""
    items = [{'key': 'vp', 'en': page['value_proposition'], 'kind': 'value proposition (one sentence)'}]
    for b in page['blocks']:
        n = b['n']
        items.append({'key': f'b{n}.title', 'en': b['title'], 'kind': f'block {n} {b["name"]} — headline'})
        items.append({'key': f'b{n}.line', 'en': b['line'], 'kind': f'block {n} {b["name"]} — one-sentence line'})
        for i, lab in enumerate(b['labels']):
            items.append({'key': f'b{n}.label{i+1}', 'en': lab, 'kind': f'block {n} {b["name"]} — chip / label'})
        items.append({'key': f'b{n}.cta', 'en': b['cta'], 'kind': f'block {n} {b["name"]} — buttons and links (raw spec incl. developer notes)'})
    return items


def cmd_inputs(story_path, home_path, work):
    story = load(story_path)
    en_pages = load(os.path.join(SITE, 'product-pages.json'))
    he_pages = load(os.path.join(SITE, 'product-pages.he.json'))
    pages = []
    for p, spec in zip(story['products'], en_pages):
        assert p['name'] == spec['name'] or p['index'] == en_pages.index(spec), (p['name'], spec['name'])
        page = {'index': p['index'], 'slug': spec['slug'], 'name': spec['name'], 'page': spec['page'],
                'value_proposition': p['value_proposition'],
                'blocks': [{'n': b['n'], 'name': b['name'], 'ladder': b['ladder'], 'title': b['title'],
                            'line': b['line'], 'labels': b['labels'], 'cta': b['cta']} for b in p['blocks']]}
        if spec['slug'] == 'exocell-mask':
            page['value_proposition'] = EXOCELL_BOARD['value_proposition']
            for b, (t, l, labs, cta) in zip(page['blocks'], EXOCELL_BOARD['blocks']):
                b.update(title=t, line=l, labels=labs, cta=cta)
            page['source'] = 'review-board-exocell.md (board version, 10/10)'
        pages.append(page)
    dump({'note': 'EN source of the v2 Story Scroll texts. EXOCELL MASK = board version.',
          'pages': pages}, os.path.join(SITE, 'storyboard.json'))

    # Home: Figma frame 69:3 text dump (y|x|path|text), de-duplicated, reading order.
    dumpd = load(home_path)
    home, seen = [], {}
    for row in dumpd['texts']:
        y, x, path, text = row.split('|', 3)
        key = 'home.' + re.sub(r'[^a-z0-9]+', '-', path.split('/')[0].lower()).strip('-')
        if text in seen:
            continue
        k = key
        i = 2
        while any(h['key'] == k for h in home):
            k = f'{key}-{i}'; i += 1
        seen[text] = k
        home.append({'key': k, 'en': text, 'kind': f'Home page, layer {path} at y={y}'})
    home.append({'key': 'nav.blocks', 'en': ' · '.join(BLOCK_NAMES),
                 'kind': 'names of the 7 story blocks (section indicator / headings in the text document)'})
    dump({'note': 'EN texts of Figma frame 69:3 "Home — Premium" (KcOGIzO76mDODxjwIrkGyo).', 'items': home},
         os.path.join(SITE, 'home.en.json'))

    os.makedirs(os.path.join(work, 'in'), exist_ok=True)
    by_slug = {p['slug']: p for p in pages}
    he_by_slug = {s['slug']: h for s, h in zip(en_pages, he_pages)}
    for gname, slugs in GROUPS:
        g = {'group': gname, 'pages': []}
        for s in slugs:
            if s == 'home':
                g['pages'].append({'slug': 'home', 'name': 'Home', 'items': home,
                                   'he_reference': {'global_nav': 'מוצרים · קטלוג · מדע · אודות · צרו קשר'}})
            else:
                p = by_slug[s]
                ref = {k: v for k, v in he_by_slug[s].items() if k not in ('index', 'page')}
                g['pages'].append({'slug': s, 'name': p['name'], 'items': page_items(p), 'he_reference': ref})
        dump(g, os.path.join(work, 'in', gname + '.json'))
    n = sum(len(page_items(p)) for p in pages) + len(home)
    print(f'pages {len(pages)}, home items {len(home)}, total items {n}')


LATIN = re.compile(r'[A-Za-z][A-Za-z0-9⁺™\-]*(?:[ .][A-Z0-9][A-Za-z0-9⁺™\-]*)*')
NUM = re.compile(r'\d+(?:[.,]\d+)?')
# Hebrew writes some small numbers as words or dual forms: שבועיים = 2 weeks, פעמיים = twice.
NUM_WORDS = {'2': ('שבועיים', 'פעמיים', 'יומיים', 'שתי', 'שני'), '3': ('תלת', 'שלוש', 'שלושה'),
             '4': ('רביעי', 'ארבע', 'ארבעה')}


def visible_en(item):
    return item.get('en_visible', item['en'])


def cmd_merge(out_path):
    out = load(out_path)
    story = load(os.path.join(SITE, 'storyboard.json'))
    home = load(os.path.join(SITE, 'home.en.json'))
    expected = {p['slug']: page_items(p) for p in story['pages']}
    expected['home'] = home['items']
    got = {}
    for g in out:
        for p in g['pages']:
            got[p['slug']] = {it['key']: it for it in p['items']}
    problems, kept = [], []
    result = {'note': 'Hebrew texts of the v2 Story Scroll site. Source: storyboard.json + home.en.json; '
                      'reviewed by a Hebrew copy editor and a fidelity checker.', 'pages': []}
    for slug, items in expected.items():
        if slug not in got:
            problems.append(f'{slug}: page missing'); continue
        rows = []
        for it in items:
            h = got[slug].get(it['key'])
            if not h:
                problems.append(f'{slug}/{it["key"]}: missing'); continue
            he = h['he'].strip()
            en_vis = h.get('en_visible', it['en']).strip() if it['key'].endswith('.cta') else it['en']
            if it['key'].endswith('.cta') and not en_vis:
                he = ''
            elif not he:
                problems.append(f'{slug}/{it["key"]}: empty Hebrew')
            if he and re.search(r'[א-ת]', he) is None and re.search(r'[a-z]{4,}', he):
                # Names kept in Latin exactly as in English (INCI, product parts, slogan) are the
                # catalogue convention; anything else is an untranslated string.
                if he.rstrip('.').lower() == en_vis.rstrip('.').lower():
                    kept.append(f'{slug}/{it["key"]}: {he}')
                else:
                    problems.append(f'{slug}/{it["key"]}: no Hebrew letters: {he}')
            for num in set(NUM.findall(en_vis)):
                if num not in he and not any(w in he for w in NUM_WORDS.get(num, ())):
                    problems.append(f'{slug}/{it["key"]}: number {num} lost: {he}')
            rows.append({'key': it['key'], 'kind': it['kind'], 'en': en_vis, 'en_spec': it['en'], 'he': he,
                         'note': h.get('note', '')})
        result['pages'].append({'slug': slug, 'name': next((p['name'] for p in story['pages'] if p['slug'] == slug), 'Home'),
                                'items': rows})
    dump(result, os.path.join(SITE, 'storyboard.he.json'))
    print(f'pages {len(result["pages"])}, items {sum(len(p["items"]) for p in result["pages"])}, '
          f'kept in Latin {len(kept)}, problems {len(problems)}')
    for p in kept:
        print('   latin:', p)
    for p in problems:
        print('   PROBLEM:', p)
    return problems


def cmd_xlsx():
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    path = os.path.join(SITE, 'mitoderm-site-blocks.xlsx')
    he = load(os.path.join(SITE, 'storyboard.he.json'))
    wb = openpyxl.load_workbook(path)
    if 'Story HE' in wb.sheetnames:
        del wb['Story HE']
    ws = wb.create_sheet('Story HE', index=wb.sheetnames.index('Storyboard') + 1 if 'Storyboard' in wb.sheetnames else None)
    ws.sheet_view.rightToLeft = True
    head = ['Page', 'Slug', 'Key', 'Field', 'עברית', 'English', 'Note']
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1E4634')
    for p in he['pages']:
        for it in p['items']:
            ws.append([p['name'], p['slug'], it['key'], it['kind'], it['he'], it['en'], it['note']])
    for col, w in zip('ABCDEFG', [24, 22, 12, 34, 70, 70, 40]):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical='top',
                                    horizontal='right' if c.column_letter == 'E' else None,
                                    readingOrder=2 if c.column_letter == 'E' else 1)
    ws.freeze_panes = 'A2'
    wb.save(path)
    print(f'Story HE: {ws.max_row - 1} rows')


if __name__ == '__main__':
    cmd = sys.argv[1]
    if cmd == 'inputs':
        cmd_inputs(*sys.argv[2:5])
    elif cmd == 'merge':
        sys.exit(1 if cmd_merge(sys.argv[2]) else 0)
    elif cmd == 'xlsx':
        cmd_xlsx()
